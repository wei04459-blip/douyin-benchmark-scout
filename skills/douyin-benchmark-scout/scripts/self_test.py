#!/usr/bin/env python3
"""Offline smoke test: workbook, classification, headers and privacy patterns."""

from __future__ import annotations

import os
import json
import importlib.util
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
ASSETS = ROOT / "assets"


def main() -> int:
    with tempfile.TemporaryDirectory(prefix="douyin-scout-test-") as tmp:
        test_root = Path(tmp)
        output = test_root / "synthetic.xlsx"
        preview = test_root / "preview"
        env = os.environ.copy()
        env.update({
            "BENCHMARK_MANIFEST": str(ASSETS / "synthetic-manifest.json"),
            "BENCHMARK_ANALYSIS": str(ASSETS / "synthetic-analysis.json"),
            "BENCHMARK_INPUT_XLSX": "",
            "BENCHMARK_OUTPUT_XLSX": str(output),
            "BENCHMARK_PREVIEW_DIR": str(preview),
        })
        subprocess.run([sys.executable, str(SCRIPTS / "build_workbook.py")], env=env, check=True)
        wb = load_workbook(output, data_only=False)
        source = wb["竞品选题分析"]
        assert source.max_column == 30
        assert source["AC2"].value == "S｜低粉高爆"
        assert wb["重点关注账号"].max_row == 16
        assert (preview / "竞品选题分析.html").is_file()

        # First-run experience must work in a completely clean data directory.
        clean_home = test_root / "clean-home"
        clean_env = os.environ.copy()
        clean_env["DOUYIN_SCOUT_HOME"] = str(clean_home)
        subprocess.run([sys.executable, str(SCRIPTS / "start.py"), "--check"], env=clean_env, check=True)
        created = json.loads((clean_home / "config.json").read_text(encoding="utf-8"))
        assert created["transcription"]["model"] == "base"
        assert created["transcription"]["threads"] == 1

        # Finalize is transactional: only now advance the cursor, delete the
        # explicit video, then rebuild Excel with the truthful post-cleanup status.
        transaction_home = test_root / "transaction-home"
        os.environ["DOUYIN_SCOUT_HOME"] = str(transaction_home)
        spec = importlib.util.spec_from_file_location("scout_run_test", SCRIPTS / "run.py")
        module = importlib.util.module_from_spec(spec)
        assert spec and spec.loader
        spec.loader.exec_module(module)
        run_dir = transaction_home / "cache" / "runs" / "test-run"
        transcript = run_dir / "transcripts" / "9000000000000000001" / "9000000000000000001.txt"
        transcript.parent.mkdir(parents=True, exist_ok=True)
        transcript.write_text("这是一段只用于离线验收的虚构口播。", encoding="utf-8")
        video = run_dir / "media" / "video.mp4"
        video.parent.mkdir(parents=True, exist_ok=True)
        video.write_bytes(b"synthetic-video-placeholder")
        manifest = json.loads((ASSETS / "synthetic-manifest.json").read_text(encoding="utf-8"))
        manifest["items"][0]["transcript_path"] = str(transcript)
        manifest["items"][0]["video_path"] = str(video)
        module.write_json(run_dir / "待分析数据.json", manifest)
        module.write_json(run_dir / module.RUN_CHECKPOINT, {"status": "awaiting_analysis", "next_keyword_cursor": 6})
        config = json.loads((ASSETS / "example-config.json").read_text(encoding="utf-8"))
        config["paths"]["workbook_python"] = sys.executable
        config = module.resolve_runtime_paths(config)
        args = SimpleNamespace(run=run_dir, analysis=ASSETS / "synthetic-analysis.json", input_workbook=None)
        module.finalize(args, config)
        assert not video.exists()
        state = module.read_json(transaction_home / "cache" / "state.json", {})
        assert state["keyword_cursor"] == 6
        assert "9000000000000000001" in state["processed_aweme_ids"]
        final_book = next(run_dir.glob("竞品选题分析_*.xlsx"))
        final_wb = load_workbook(final_book)
        assert final_wb["竞品选题分析"]["Z2"].value == "已完成转录与分析，原视频已删除"

        outside_video = test_root / "must-not-delete.mp4"
        outside_video.write_bytes(b"protected")
        guard_run = transaction_home / "cache" / "runs" / "guard-run"
        guard_transcript = guard_run / "transcript.txt"
        guard_transcript.parent.mkdir(parents=True, exist_ok=True)
        guard_transcript.write_text("guard", encoding="utf-8")
        guard_manifest = json.loads((ASSETS / "synthetic-manifest.json").read_text(encoding="utf-8"))
        guard_manifest["items"][0].update({"transcript_path": str(guard_transcript), "video_path": str(outside_video)})
        module.write_json(guard_run / "待分析数据.json", guard_manifest)
        try:
            module.delete_processed_videos(guard_run / "待分析数据.json", ASSETS / "synthetic-analysis.json", final_book)
            raise AssertionError("cleanup accepted a video outside the run directory")
        except RuntimeError as exc:
            assert "运行目录之外" in str(exc)
        assert outside_video.is_file()

        # Legacy migration merges dedupe state and copies only safe artifacts.
        legacy = test_root / "legacy"
        (legacy / "runs" / "old").mkdir(parents=True)
        (legacy / "state.json").write_text(json.dumps({"keyword_cursor": 3, "seen_aweme_ids": ["1"], "processed_aweme_ids": ["1"]}), encoding="utf-8")
        shutil.copy2(output, legacy / "runs" / "old" / "old.xlsx")
        migrate_home = test_root / "migrated"
        subprocess.run(
            [sys.executable, str(SCRIPTS / "migrate_legacy.py"), "--source", str(legacy), "--destination", str(migrate_home), "--apply"],
            check=True,
        )
        migrated = json.loads((migrate_home / "cache" / "state.json").read_text(encoding="utf-8"))
        assert migrated["seen_aweme_ids"] == ["1"]
        assert (migrate_home / "竞品选题分析.xlsx").is_file()
    forbidden = [
        "/Users/" + "yuwan",
        "yuwan" + "-knowledge-base",
        "api" + "-business",
        "BEGIN " + "PRIVATE KEY",
        "gh" + "p_",
    ]
    for path in ROOT.rglob("*"):
        if path.is_file() and path.suffix.lower() in {".py", ".md", ".json", ".yaml", ".yml", ".mjs", ".command", ".bat"}:
            text = path.read_text(encoding="utf-8", errors="ignore")
            matches = [token for token in forbidden if token in text]
            if matches:
                raise AssertionError(f"Privacy pattern in {path}: {matches}")
    print("SELF_TEST_OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
