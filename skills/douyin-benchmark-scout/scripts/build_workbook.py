#!/usr/bin/env python3
"""Create or refresh the competitor workbook with openpyxl only."""

from __future__ import annotations

import html
import json
import os
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADERS = [
    "序号", "内容方向/赛道", "视频链接", "账号名", "账号粉丝数", "账号总获赞",
    "视频发布时间", "视频时长（秒）", "视频点赞", "视频评论", "视频收藏", "视频转发",
    "收藏/点赞比", "视频标题", "内容关键词", "选题类型", "选题概括", "开头钩子/模式",
    "内容结构", "情绪设计", "内容摘要与事实边界", "爆款归因", "账号可迁移思路",
    "命中搜索关键词", "完整口播稿路径", "本地视频状态", "赞粉比", "互动/粉丝比",
    "爆款等级", "判断依据",
]
REQUIRED_ANALYSIS = ["track", "type", "topic_summary", "hook", "structure", "emotion", "summary", "viral", "migration"]
THIN = Side(style="thin", color="B4C6E7")
LEVEL_FILL = {"S": "FCE4D6", "A": "FFF2CC", "B": "FFFFFF", "C": "E7E6E6", "待": "F4CCCC"}


def read_json(name: str) -> dict:
    return json.loads(Path(os.environ[name]).read_text(encoding="utf-8"))


def number(value):
    try:
        return int(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def normalize(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\\n", " ")).strip()


def classify(fans, likes, rules):
    if not fans or fans <= 0:
        return "待补粉丝数", "缺少公开粉丝数，暂不判定"
    likes = likes or 0
    ratio = likes / fans
    if fans <= rules["s_max_fans"] and likes >= rules["s_min_likes"] and ratio >= rules["s_min_like_fan_ratio"]:
        return "S｜低粉高爆", "粉丝≤10万、点赞≥1万且赞粉比≥100%"
    if fans <= rules["a_max_fans"] and likes >= rules["a_min_likes"] and ratio >= rules["a_min_like_fan_ratio"]:
        return "A｜重点观察", "粉丝≤30万、点赞≥3万且赞粉比≥30%"
    if fans >= rules["big_account_fans"] and ratio < rules["big_account_ratio_ceiling"]:
        return "C｜大号常规流量", "粉丝≥100万且赞粉比<10%，更可能由账号基本盘驱动"
    return "B｜常规样本", "未达到重点阈值"


def make_row(item, analysis, serial, rules):
    fans = number(item.get("follower_count"))
    likes = number(item.get("liked_count")) or 0
    comments = number(item.get("comment_count")) or 0
    favorites = number(item.get("collected_count")) or 0
    shares = number(item.get("share_count")) or 0
    level, rationale = classify(fans, likes, rules)
    published = None
    if number(item.get("create_time")):
        published = datetime.fromtimestamp(number(item["create_time"]))
    title = normalize(item.get("title") or item.get("desc"))
    topics = normalize(analysis.get("topics")) or " ".join(re.findall(r"#[^\s#]+", title))
    source_keywords = item.get("source_keywords") or item.get("source_keyword") or ""
    if isinstance(source_keywords, list):
        source_keywords = "、".join(map(str, source_keywords))
    transcript = str(item.get("transcript_path") or "")
    video_status = str(item.get("video_status") or "")
    if not video_status:
        if item.get("video_deleted_at"):
            video_status = "已完成转录与分析，原视频已删除"
        elif transcript and item.get("video_path"):
            video_status = "已完成转录与分析，等待安全清理"
        elif transcript:
            video_status = "已保留口播稿，本地原视频不存在"
        else:
            video_status = str(item.get("video_path") or "未下载")
    return [
        serial, analysis["track"], item.get("aweme_url") or f"https://www.douyin.com/video/{item['aweme_id']}",
        item.get("nickname") or "", fans, number(item.get("creator_total_favorited")), published,
        number(item.get("duration_seconds")), likes, comments, favorites, shares,
        favorites / likes if likes else None, title, topics, analysis["type"], analysis["topic_summary"],
        analysis["hook"], analysis["structure"], analysis["emotion"], analysis["summary"], analysis["viral"],
        analysis["migration"], str(source_keywords), transcript, video_status,
        likes / fans if fans else None, (likes + comments + favorites + shares) / fans if fans else None,
        level, rationale,
    ]


def style_source(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    widths = [8, 18, 38, 16, 13, 14, 13, 12, 12, 12, 12, 12, 13, 42, 24, 18, 32, 34, 40, 28, 46, 40, 46, 20, 38, 32, 13, 15, 18, 42]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in range(2, ws.max_row + 1):
        base = "DDEBF7" if row % 2 == 0 else "FFFFFF"
        level = str(ws.cell(row, 29).value or "")
        for cell in ws[row]:
            cell.fill = PatternFill("solid", fgColor=base)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row, 29).fill = PatternFill("solid", fgColor=LEVEL_FILL.get(level[:1], "FFF2CC"))
        ws.cell(row, 29).font = Font(bold=True)
        ws.row_dimensions[row].height = 78
        ws.cell(row, 7).number_format = "yyyy-mm-dd"
        for col in (13, 27, 28):
            ws.cell(row, col).number_format = "0.0%"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:AD{max(1, ws.max_row)}"


def rebuild_focus(wb, source, rules):
    if "重点关注账号" in wb.sheetnames:
        del wb["重点关注账号"]
    ws = wb.create_sheet("重点关注账号")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:M1")
    ws["A1"] = "低粉高爆 · 重点关注账号"
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:M2")
    ws["A2"] = "核心原则：先看单条视频是否突破账号基本盘，再看绝对互动量；同一账号保留赞粉比最高的视频。"
    ws["A2"].fill = PatternFill("solid", fgColor="D9EAF7")
    params = [
        ("S 最大粉丝数", rules["s_max_fans"]), ("S 最低点赞数", rules["s_min_likes"]),
        ("S 最低赞粉比", rules["s_min_like_fan_ratio"]), ("A 最大粉丝数", rules["a_max_fans"]),
        ("A 最低点赞数", rules["a_min_likes"]), ("A 最低赞粉比", rules["a_min_like_fan_ratio"]),
        ("大号粉丝门槛", rules["big_account_fans"]), ("大号常规赞粉比上限", rules["big_account_ratio_ceiling"]),
    ]
    ws["A4"] = "判定参数"
    ws["B4"] = "当前值"
    for row_index, (name, value) in enumerate(params, 5):
        ws.cell(row_index, 1, name)
        ws.cell(row_index, 2, value)
    headers = ["关注级别", "账号名", "粉丝数", "代表视频", "点赞", "收藏", "评论", "转发", "赞粉比", "互动/粉丝比", "发布时间", "视频链接", "判断依据"]
    for col, value in enumerate(headers, 1):
        ws.cell(15, col, value)
    best = {}
    for row in source.iter_rows(min_row=2, values_only=True):
        level, account = str(row[28] or ""), str(row[3] or "").strip()
        if not account or not level.startswith(("S", "A")):
            continue
        ratio = row[26] or 0
        if account not in best or ratio > best[account][0]:
            best[account] = (ratio, row)
    ranked = sorted(best.values(), key=lambda x: (0 if str(x[1][28]).startswith("S") else 1, -x[0]))
    for _, row in ranked:
        ws.append([row[28], row[3], row[4], row[13], row[8], row[10], row[9], row[11], row[26], row[27], row[6], row[2], row[29]])
        fill = LEVEL_FILL.get(str(row[28])[:1], "FFF2CC")
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws[15]:
        cell.fill = PatternFill("solid", fgColor="2F5496")
        cell.font = Font(color="FFFFFF", bold=True)
    if ws.max_row >= 16:
        table = Table(displayName="重点关注账号表", ref=f"A15:M{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    for col, width in {"A":16,"B":16,"C":12,"D":38,"E":11,"F":11,"G":11,"H":11,"I":14,"J":15,"K":13,"L":38,"M":42}.items():
        ws.column_dimensions[col].width = width
    for row in range(16, ws.max_row + 1):
        ws.cell(row, 9).number_format = ws.cell(row, 10).number_format = "0.0%"
        ws.cell(row, 11).number_format = "yyyy-mm-dd"
        ws.row_dimensions[row].height = 58
    ws.freeze_panes = "A16"


def write_preview(ws, path: Path, limit=12):
    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, limit), values_only=True))
    body = ["<meta charset='utf-8'><style>table{border-collapse:collapse;font:12px sans-serif}th,td{border:1px solid #b4c6e7;padding:6px;max-width:260px}th{background:#1f4e78;color:white}tr:nth-child(even){background:#ddebf7}</style><table>"]
    for ridx, row in enumerate(rows):
        tag = "th" if ridx == 0 else "td"
        body.append("<tr>" + "".join(f"<{tag}>{html.escape(str(v or ''))}</{tag}>" for v in row) + "</tr>")
    body.append("</table>")
    path.write_text("\n".join(body), encoding="utf-8")


def main():
    manifest = read_json("BENCHMARK_MANIFEST")
    analysis_payload = read_json("BENCHMARK_ANALYSIS")
    analysis_by_id = analysis_payload.get("items", analysis_payload)
    items = manifest.get("items", [])
    for item in items:
        aid = str(item.get("aweme_id"))
        analysis = analysis_by_id.get(aid)
        if not isinstance(analysis, dict):
            raise RuntimeError(f"Missing analysis for {aid}")
        missing = [field for field in REQUIRED_ANALYSIS if not normalize(analysis.get(field))]
        if missing:
            raise RuntimeError(f"Analysis {aid} is incomplete: {', '.join(missing)}")
    input_path = Path(os.environ.get("BENCHMARK_INPUT_XLSX", ""))
    wb = load_workbook(input_path) if input_path.is_file() else Workbook()
    if "竞品选题分析" in wb.sheetnames:
        ws = wb["竞品选题分析"]
    else:
        ws = wb.active
        ws.title = "竞品选题分析"
    for idx, header in enumerate(HEADERS, 1):
        ws.cell(1, idx, header)
    existing = {}
    for row in range(2, ws.max_row + 1):
        match = re.search(r"video/(\d+)", str(ws.cell(row, 3).value or ""))
        if match:
            existing[match.group(1)] = row
    rules = {"s_max_fans":100000,"s_min_likes":10000,"s_min_like_fan_ratio":1.0,"a_max_fans":300000,"a_min_likes":30000,"a_min_like_fan_ratio":0.3,"big_account_fans":1000000,"big_account_ratio_ceiling":0.1}
    rules.update(manifest.get("rules") or {})
    next_row = max(2, ws.max_row + 1)
    for item in items:
        aid = str(item["aweme_id"])
        target = existing.get(aid, next_row)
        if aid not in existing:
            next_row += 1
        values = make_row(item, analysis_by_id[aid], target - 1, rules)
        for col, value in enumerate(values, 1):
            ws.cell(target, col, value)
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 1, row - 1)
    style_source(ws)
    rebuild_focus(wb, ws, rules)
    output = Path(os.environ["BENCHMARK_OUTPUT_XLSX"])
    output.parent.mkdir(parents=True, exist_ok=True)
    temp = output.with_suffix(".tmp.xlsx")
    wb.save(temp)
    check = load_workbook(temp, read_only=False, data_only=False)
    if [check["竞品选题分析"].cell(1, i).value for i in range(1, 31)] != HEADERS:
        raise RuntimeError("Workbook header contract validation failed")
    for sheet in check.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(err in cell.value for err in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")):
                    raise RuntimeError(f"Formula error in {sheet.title}!{cell.coordinate}")
    temp.replace(output)
    preview_dir = Path(os.environ["BENCHMARK_PREVIEW_DIR"])
    preview_dir.mkdir(parents=True, exist_ok=True)
    write_preview(check["竞品选题分析"], preview_dir / "竞品选题分析.html")
    write_preview(check["重点关注账号"], preview_dir / "重点关注账号.html", limit=30)
    print(json.dumps({"output": str(output), "rows": check["竞品选题分析"].max_row - 1, "focus_accounts": max(0, check["重点关注账号"].max_row - 15)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
